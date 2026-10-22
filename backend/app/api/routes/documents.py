"""Document management endpoints."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import UUID

from docx import Document as DocxDocument
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from pypdf import PdfReader

from app.api.dependencies import get_document_service, get_settings
from app.api.schemas.documents import (
    DocumentDetailResponse,
    DocumentResponse,
    DocumentTagLiteral,
    DocumentUpsertRequest,
)
from app.core.config import Settings
from app.domain.models.document import DocumentRecord, DocumentTag
from app.ingestion.splitters import PAGE_BREAK
from app.services.document_service import DocumentService

router = APIRouter()

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".md"}


@router.get("/", response_model=list[DocumentResponse])
def list_documents(
    document_service: DocumentService = Depends(get_document_service),
) -> list[DocumentResponse]:
    documents = document_service.list_documents()
    return [to_response(doc) for doc in documents]


@router.post("/", response_model=DocumentResponse, status_code=status.HTTP_201_CREATED)
async def upload_document(
    file: UploadFile = File(...),
    filename: str | None = Form(None),
    tag: DocumentTagLiteral = Form("other"),
    document_service: DocumentService = Depends(get_document_service),
    settings: Settings = Depends(get_settings),
) -> DocumentResponse:
    display_name = (filename or file.filename or "").strip()
    if not display_name:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Filename is required.")
    try:
        content = await extract_text_from_upload(file, max_bytes=settings.max_upload_mb * 1024 * 1024)
    except UploadTooLargeError as exc:
        raise HTTPException(status_code=413, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to process document.",
        ) from exc
    finally:
        await file.close()

    try:
        record = document_service.upsert_document(filename=display_name, content=content, tag=DocumentTag(tag))
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc

    return to_response(record)


@router.put("/{document_id}", response_model=DocumentResponse)
def update_document(
    document_id: UUID,
    payload: DocumentUpsertRequest,
    document_service: DocumentService = Depends(get_document_service),
) -> DocumentResponse:
    record = document_service.get_document(document_id)
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    try:
        updated = document_service.upsert_document(
            filename=payload.filename,
            content=payload.content,
            document_id=document_id,
            tag=DocumentTag(payload.tag),
        )
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    return to_response(updated)


@router.get("/{document_id}", response_model=DocumentDetailResponse)
def get_document(
    document_id: UUID,
    document_service: DocumentService = Depends(get_document_service),
) -> DocumentDetailResponse:
    result = document_service.get_document_with_content(document_id)
    if not result:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    record, content = result
    return to_response(record, content=content)


@router.delete("/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_document(
    document_id: UUID,
    document_service: DocumentService = Depends(get_document_service),
) -> None:
    record = document_service.get_document(document_id)
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    document_service.delete_document(document_id)


def to_response(record: DocumentRecord, *, content: str | None = None) -> DocumentResponse | DocumentDetailResponse:
    payload = {
        "id": record.id,
        "filename": record.filename,
        "summary": record.summary,
        "tag": record.tag.value,
        "chunk_count": record.chunk_count,
        "uploaded_at": record.uploaded_at,
        "updated_at": record.updated_at,
    }
    if content is not None:
        return DocumentDetailResponse(**payload, content=content)
    return DocumentResponse(**payload)


class UploadTooLargeError(ValueError):
    pass


async def extract_text_from_upload(upload: UploadFile, *, max_bytes: int | None = None) -> str:
    extension = Path(upload.filename or "").suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported file type. Allowed extensions: .txt, .pdf, .docx, .md.")

    chunks: list[bytes] = []
    size = 0
    while chunk := await upload.read(1 << 20):
        size += len(chunk)
        if max_bytes is not None and size > max_bytes:
            raise UploadTooLargeError(f"File is larger than {max_bytes // (1024 * 1024)} MB.")
        chunks.append(chunk)
    raw_bytes = b"".join(chunks)
    if not raw_bytes:
        raise ValueError("Uploaded file is empty.")

    if extension in {".txt", ".md"}:
        try:
            text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = raw_bytes.decode("utf-8", errors="ignore")
        content = text.strip()
    elif extension == ".pdf":
        reader = PdfReader(BytesIO(raw_bytes))
        # Keep page boundaries (form feed) so chunks can cite page numbers. v1
        # joined pages with blank lines and lost them, so citations never had pages.
        pages = [(page.extract_text() or "").replace(PAGE_BREAK, " ").strip() for page in reader.pages]
        content = PAGE_BREAK.join(pages)  # no strip(): it would eat leading empty pages and shift numbers
    else:  # .docx
        document = DocxDocument(BytesIO(raw_bytes))
        segments: list[str] = []
        segments.extend(p.text.strip() for p in document.paragraphs if p.text.strip())
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    segments.append("\t".join(cells))
        content = "\n\n".join(segments).strip()

    if not content.strip():
        raise ValueError("No readable text content found in the uploaded file.")
    return content
